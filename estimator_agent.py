"""CIVINITY photo-to-estimate agent. Research results are evidence, not automatically verified shop prices."""
import base64, io, json, os, re, threading, uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from flask import Blueprint, jsonify, request, send_file
from openai import OpenAI
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

bp=Blueprint('estimator_agent',__name__)
_jobs={}; _lock=threading.Lock(); _pool=ThreadPoolExecutor(max_workers=2)

def _json(s):
    s=s.strip(); a=s.find('{'); b=s.rfind('}')
    if a<0 or b<a: raise ValueError('AI neatgrieza JSON objektu')
    return json.loads(s[a:b+1])

def _client():
    return OpenAI(api_key=os.environ['OPENAI_API_KEY'],timeout=110,max_retries=1)

def _update(jid,**kw):
    with _lock: _jobs[jid].update(kw)

def _run(jid,description,images):
    try:
        client=_client(); model=os.getenv('ESTIMATOR_MODEL','gpt-4.1')
        _update(jid,status='vision',progress=12)
        content=[{'type':'text','text':description}]
        for pic in images:
            content.append({'type':'image_url','image_url':{'url':pic,'detail':'high'}})
        spec_prompt='''Tu esi pieredzējis Latvijas būvdarbu tāmētājs. Analizē lietotāja tekstu, shēmas un foto. Neizdomā izmērus, ko nevar nolasīt; izskaidro ģeometriju un atzīmē pieņēmumus. Ja atkritumu laukums atrodas kopējā laukumā, nepieskaiti to otrreiz. Izveido pilnu būvdarbu tehnoloģiju, no demontāžas līdz nodošanai, atsevišķi materiālus, darbus, mehānismus, izvešanas, rezerves, zudumus un drenāžas riskus. Neizdomā veikalu cenas. Atgriez tikai JSON: {"title":"...","summary":"...","geometry":[{"name":"...","formula":"...","area_m2":0}],"assumptions":["..."],"clarifications":["..."],"work_scope":["..."],"items":[{"name":"...","category":"labor|material|machine","unit":"m2|m3|m|gab.|kg|diena|kompl.","quantity":1,"technical_spec":"...","quantity_formula":"...","waste_pct":0}],"research_queries":["..."],"risks":["..."]}. Izmanto latviešu valodu. Precīzi ievēro lietotāja norādītos izmērus un pārbaudi aritmētiku.'''
        res=client.chat.completions.create(model=model,messages=[{'role':'system','content':spec_prompt},{'role':'user','content':content}],response_format={'type':'json_object'},max_tokens=8500)
        spec=_json(res.choices[0].message.content)
        _update(jid,status='research',progress=37,spec=spec)
        queries=(spec.get('research_queries') or [])[:14]
        itemnames=', '.join(str(x.get('name','')) for x in (spec.get('items') or [])[:38])
        research_prompt='''Meklē Latvijas internetā AKTUĀLUS publicētus būvmateriālu piedāvājumus un būvdarbu cenrāžus konkrētajam projektam. Izmanto web search. Atgriez tikai JSON ar {"offers":[{"item":"tāmes pozīcija","supplier":"...","product":"...","url":"https://...","price_eur":12.34,"unit":"m2","vat_included":true,"notes":"...","evidence":"published_page|search_snippet|price_range"}],"research_notes":["..."]}. Nedrīkst izdomāt cenas vai saites. Ja cenas nav, NEVEIDO piedāvājumu. Atšķir darbu likmes no materiālu cenām. Nekādā gadījumā neizmanto AI izdomātus cenu avotus. Preferē Latvijas tirgotājus un būvniekus. Maksimums 45 piedāvājumi. Lietotāja uzdevums: '''+description[:2500]+'\nPozīcijas: '+itemnames[:4200]+'\nMeklējumi: '+json.dumps(queries,ensure_ascii=False)
        evidence={'offers':[],'research_notes':[]}
        try:
            rr=client.responses.create(model=os.getenv('ESTIMATOR_RESEARCH_MODEL','gpt-4.1'),tools=[{'type':'web_search_preview','search_context_size':'high'}],input=research_prompt,max_output_tokens=5500,timeout=110)
            evidence=_json(rr.output_text)
        except Exception as e:
            evidence['research_notes']=['Tiešsaistes izpēte nav pabeigta: '+str(e)[:180]]
        # An LLM web search result alone is never a shop-page verification.
        cleaned=[]
        for o in evidence.get('offers',[])[:50]:
            if not isinstance(o,dict):continue
            url=str(o.get('url',''))
            if not url.startswith('https://') or len(url)>900:continue
            try:
                price=float(o.get('price_eur'))
                if not 0<price<1e7:continue
            except (ValueError,TypeError):continue
            o['price_eur']=round(price,2)
            o['verification']='web_search_reported_not_independently_verified'
            cleaned.append(o)
        evidence['offers']=cleaned
        _update(jid,status='estimate',progress=70,evidence=evidence)
        prompt='''Sagatavo PROFESIONĀLU Latvijas būvdarbu TĀMI pēc tehniskās specifikācijas un interneta cenu izpētes. Neizmanto nulles cenas kā bezmaksas materiālus. Katram darbam un materiālam, kam nav pierādītas cenas, drīksti ieteikt BUDŽETA NOVĒRTĒJUMU, taču obligāti atzīmē statusu "provisional" un paskaidro, ka tā nav pārbaudīta tirgus cena. Ja cena ir salīdzināma ar publicēto piedāvājumu, status "web_researched" un norādi avota URL. Atsevišķi labor, material, machine izmaksas katrā rindā, bez dubultskaitīšanas. Aprēķini precīzus apjomus, iekļauj zudumus un transportu, demontāžu, atkritumus, materiālus, darbus, mehānismus, apmales, nepieciešamo grunts sagatavošanu un objekta nodošanu. Sadalījums pa darbu posmiem. Atkritumu konteineru zona atsevišķā nodaļā, bet nedubulto kopējo seguma laukumu. PVN 21%. Neuzdod izdomātas cenas kā veikalu cenas. Atgriez tikai JSON: {"title":"...","summary":"...","geometry":[...],"assumptions":["..."],"clarifications":["..."],"sections":[{"name":"...","items":[{"name":"...","unit":"m2","quantity":1,"labor":0,"material":0,"machine":0,"status":"web_researched|provisional|missing","source_url":"...","technical_spec":"...","quantity_formula":"...","note":"..."}]}],"risks":["..."],"methodology":"..."}. Visi skaitļi EUR bez PVN. Ja avota cenā PVN iekļauts, pārrēķini bez PVN. Neizmanto cenu, ja nav skaidra mērvienība vai specifikācija. Vismaz 15 detalizētas pozīcijas, ja apjoms to prasa.'''
        resp=client.chat.completions.create(model=model,messages=[{'role':'system','content':prompt},{'role':'user','content':json.dumps({'specification':spec,'market_research':evidence},ensure_ascii=False)}],response_format={'type':'json_object'},max_tokens=11000)
        result=_json(resp.choices[0].message.content)
        total=0; missing=0
        for section in result.get('sections',[]):
            for item in section.get('items',[]):
                def numeric(k):
                    try: v=float(item.get(k) or 0); return round(max(0,min(v,1e8)),4)
                    except (TypeError,ValueError):return 0
                for k in ('quantity','labor','material','machine'):item[k]=numeric(k)
                item['subtotal']=round(item['quantity']*(item['labor']+item['material']+item['machine']),2)
                total+=item['subtotal']
                if item.get('status')=='missing':missing+=1
        result['subtotal']=round(total,2);result['vat']=round(total*.21,2);result['total']=round(total*1.21,2)
        result['market_offers']=cleaned;result['research_notes']=evidence.get('research_notes',[])
        result['disclaimer']='Provizoriska tāme: web_search atrastās cenas nav automātiski apstiprināti veikalu piedāvājumi; status provisional nav pārbaudīta tirgus cena. Pirms piedāvājuma apstiprināt specifikācijas, PVN, cenas, piegādi un apjomus.'
        result['created_at']=datetime.now(timezone.utc).isoformat()
        _update(jid,status='done',progress=100,result=result)
    except Exception as e:
        _update(jid,status='failed',error=str(e)[:350],progress=100)

@bp.post('/api/agent/estimate')
def start():
    if not os.getenv('OPENAI_API_KEY'):return jsonify(error='Render vidē nav OPENAI_API_KEY'),503
    body=request.get_json(silent=True) or {}
    description=str(body.get('description',''))[:12000]
    images=body.get('images') or []
    if len(description.strip())<10:return jsonify(error='Nepieciešams darbu apraksts'),400
    if not isinstance(images,list) or len(images)>6:return jsonify(error='Maksimums 6 attēli'),400
    for img in images:
        if not isinstance(img,str) or not re.match(r'^data:image/(jpeg|png|webp);base64,',img) or len(img)>4500000:
            return jsonify(error='Nepareizs foto formāts vai izmērs'),400
    jid=uuid.uuid4().hex
    with _lock:
        if len(_jobs)>150:
            for k in list(_jobs)[:60]:_jobs.pop(k,None)
        _jobs[jid]={'status':'queued','progress':0}
    _pool.submit(_run,jid,description,images)
    return jsonify(job_id=jid),202

@bp.get('/api/agent/estimate/<jid>')
def get_job(jid):
    with _lock:job=_jobs.get(jid)
    return (jsonify(job) if job else (jsonify(error='Uzdevums nav atrasts; serveris iespējams restartējās'),404))

@bp.post('/api/agent/xlsx')
def export_xlsx():
    data=request.get_json(silent=True) or {}; r=data.get('result') or {}; c=data.get('company') or {}
    if not isinstance(r.get('sections'),list):return jsonify(error='Tāme nav pieejama'),400
    wb=Workbook();ws=wb.active;ws.title='Tāme'
    ws.merge_cells('A1:H1');ws['A1']=str(c.get('name') or 'CIVINITY')[:180];ws['A1'].font=Font(size=19,bold=True,color='0A3960')
    ws.merge_cells('A2:H2');ws['A2']=' | '.join(str(c.get(k) or '') for k in ('reg','address','iban','bank'))[:400]
    ws.merge_cells('A4:H4');ws['A4']=str(r.get('title') or 'BŪVDARBU TĀME');ws['A4'].font=Font(size=15,bold=True,color='0A3960')
    ws.merge_cells('A5:H5');ws['A5']=str(data.get('object') or '')
    ws.append([]);ws.append(['Nr.','Pozīcija / specifikācija','Mērv.','Daudz.','Darbs €/v.','Materiāli €/v.','Mehānismi €/v.','Kopā bez PVN €'])
    for cell in ws[7]:cell.fill=PatternFill('solid',fgColor='143D61');cell.font=Font(color='FFFFFF',bold=True)
    n=0
    for sec in r.get('sections',[]):
        ws.append(['',str(sec.get('name',''))]);ws.cell(ws.max_row,2).font=Font(bold=True,color='0A3960')
        for i in sec.get('items',[]):
            n+=1; row=ws.max_row+1
            ws.append([n,str(i.get('name',''))+'\n'+str(i.get('technical_spec',''))+'\n'+str(i.get('note','')),i.get('unit'),i.get('quantity'),i.get('labor'),i.get('material'),i.get('machine'),f'=D{row}*(E{row}+F{row}+G{row})'])
            ws.cell(row,2).alignment=Alignment(wrap_text=True,vertical='top');ws.row_dimensions[row].height=43
    last=ws.max_row
    ws.append(['','KOPĀ BEZ PVN','','','','','',f'=SUM(H8:H{last})'])
    ws.append(['','PVN 21%','','','','',f'=H{last+1}*0.21'])
    ws.append(['','KOPĀ AR PVN','','','','',f'=H{last+1}+H{last+2}'])
    ws2=wb.create_sheet('Cenu avoti');ws2.append(['Pozīcija','Piegādātājs','Prece','Cena EUR','Mērvienība','PVN iekļauts','Saite','Pārbaudes statuss'])
    for o in r.get('market_offers',[]):ws2.append([o.get('item'),o.get('supplier'),o.get('product'),o.get('price_eur'),o.get('unit'),o.get('vat_included'),o.get('url'),o.get('verification')])
    ws3=wb.create_sheet('Pieņēmumi un riski')
    for heading,key in [('Ģeometrija','geometry'),('Pieņēmumi','assumptions'),('Jāprecizē','clarifications'),('Riski','risks'),('Izpētes piezīmes','research_notes')]:
        ws3.append([heading]);ws3.cell(ws3.max_row,1).font=Font(bold=True,color='0A3960')
        for v in r.get(key,[]):ws3.append([json.dumps(v,ensure_ascii=False) if isinstance(v,dict) else str(v)])
    ws3.append([r.get('disclaimer','')])
    for sh in wb:
        sh.freeze_panes='C8' if sh==ws else 'A2'
        sh.sheet_view.showGridLines=False
        if sh==ws:
            for col,w in {'A':7,'B':70,'C':12,'D':12,'E':15,'F':17,'G':17,'H':21}.items():sh.column_dimensions[col].width=w
        elif sh==ws2:
            for col in 'ABCDEFGH':sh.column_dimensions[col].width=24 if col!='G' else 75
        else:sh.column_dimensions['A'].width=110
    # User logo is supplied by the existing settings form, not hard-coded.
    logo=c.get('logo','')
    if isinstance(logo,str) and logo.startswith('data:image/'):
        try:
            from openpyxl.drawing.image import Image
            binary=base64.b64decode(logo.split(',',1)[1]); im=Image(io.BytesIO(binary));im.width=min(im.width,145);im.height=min(im.height,58);ws.add_image(im,'G1')
        except Exception:pass
    buf=io.BytesIO();wb.save(buf);buf.seek(0)
    return send_file(buf,mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',as_attachment=True,download_name='CIVINITY_AI_Tame_Agent.xlsx')
